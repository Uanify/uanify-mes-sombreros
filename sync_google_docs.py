import os
import json
import argparse
from pathlib import Path
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaFileUpload

SCOPES = [
    'https://www.googleapis.com/auth/drive',
    'https://www.googleapis.com/auth/documents'
]

CLIENT_SECRET_FILE = Path(__file__).parent / 'client_secret.json'
TOKEN_FILE = Path(__file__).parent / 'token.json'
UANIFY_LOGO_PATH = Path(__file__).parent / 'assets' / 'uanify_brand_logo.png'

# --- PALETA CORPORATIVA OFICIAL UANIFY (Estilo Capturas) ---
COLOR_PRIMARY_DARK = RGBColor(15, 23, 42)    # #0F172A Slate 900 (Títulos principales H1, H2)
COLOR_BRAND_BLUE   = RGBColor(37, 99, 235)   # #2563EB Blue 600 (Subtítulos, H3, Acentos de Marca)
COLOR_BODY         = RGBColor(51, 65, 85)    # #334155 Slate 700 (Texto corrido claro y legible)
COLOR_MUTED        = RGBColor(100, 116, 139) # #64748B Slate 500 (Metadatos, pie de página)
COLOR_CARD_BORDER  = "2563EB"                # Azul Uanify para borde lateral de notas/callouts
COLOR_LINE_BORDER  = "CBD5E1"                # Slate 300 para separadores horizontales de sección
COLOR_BG_CARD      = "F8FAFC"                # Slate 50 para fondos de notas / filas alternadas
COLOR_TH_BG        = "F1F5F9"                # Slate 100 para encabezados de tabla

def clean_emojis(text: str) -> str:
    """Elimina emojis y símbolos pictográficos para garantizar una apariencia ejecutiva 100% limpia."""
    if not text:
        return ""
    # Rango de emojis y caracteres decorativos Unicode
    emoji_pattern = re.compile(
        "[\U00010000-\U0010ffff"
        "\u2600-\u26ff"
        "\u2700-\u27bf"
        "\ufe0f"
        "\u200d"
        "\u2300-\u23ff"
        "\u2b50-\u2b55"
        "\u203c-\u2049"
        "\u25aa-\u25fe]",
        flags=re.UNICODE
    )
    cleaned = emoji_pattern.sub('', text)
    # Limpiar dobles espacios residuales tras remover emojis
    return re.sub(r' {2,}', ' ', cleaned).strip()

def get_credentials():
    if not TOKEN_FILE.exists():
        raise FileNotFoundError("token.json no existe. Ejecuta exchange_token.py primero.")

    with open(TOKEN_FILE, 'r', encoding='utf-8') as f:
        data = json.load(f)

    with open(CLIENT_SECRET_FILE, 'r', encoding='utf-8') as f:
        cs = json.load(f)['installed']

    creds = Credentials(
        token=data.get('access_token'),
        refresh_token=data.get('refresh_token'),
        token_uri="https://oauth2.googleapis.com/token",
        client_id=cs['client_id'],
        client_secret=cs['client_secret'],
        scopes=SCOPES
    )

    if creds.expired and creds.refresh_token:
        creds.refresh(Request())
        data['access_token'] = creds.token
        with open(TOKEN_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f)

    return creds

def get_services():
    creds = get_credentials()
    drive = build('drive', 'v3', credentials=creds)
    docs = build('docs', 'v1', credentials=creds)
    return drive, docs

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_card_borders(cell, border_color="2563EB"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_horizontal_divider(doc, space_after=10):
    """Agrega una línea divisoria horizontal sutil como en las plantillas oficiales de Uanify."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(space_after)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="8" w:color="{COLOR_LINE_BORDER}"/></w:pBdr>')
    p._p.get_or_add_pPr().append(pBdr)

def add_formatted_runs(p, text, default_color=COLOR_BODY, default_size=10, default_font="Arial", default_bold=False):
    """Parsea markdown inline (**negrita**, *cursiva* y `codigo`) dentro de cualquier parrafo o celda, sin emojis."""
    clean_text = clean_emojis(text)
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', clean_text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            r = p.add_run(part[2:-2])
            r.bold = True
            r.font.name = default_font
            r.font.size = Pt(default_size)
            r.font.color.rgb = COLOR_PRIMARY_DARK
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            r = p.add_run(part[1:-1])
            r.italic = True
            r.bold = default_bold
            r.font.name = default_font
            r.font.size = Pt(default_size)
            r.font.color.rgb = default_color
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            r = p.add_run(part[1:-1])
            r.bold = True
            r.font.name = "Consolas"
            r.font.size = Pt(default_size - 0.5)
            r.font.color.rgb = COLOR_BRAND_BLUE
        else:
            r = p.add_run(part)
            r.bold = default_bold
            r.font.name = default_font
            r.font.size = Pt(default_size)
            r.font.color.rgb = default_color

def md_to_docx(md_path, docx_path, doc_title):
    doc = docx.Document()
    
    # 1. Configuración de Márgenes Ejecutivos y Encabezados
    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.85)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
        # Encabezado Oficial con Logotipo Uanify alineado a la derecha
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if UANIFY_LOGO_PATH.exists():
            h_run = header_p.add_run()
            h_run.add_picture(str(UANIFY_LOGO_PATH), width=Inches(1.2))
        else:
            h_run = header_p.add_run("uanify")
            h_run.font.name = "Arial"
            h_run.font.size = Pt(13)
            h_run.bold = True
            h_run.font.color.rgb = COLOR_PRIMARY_DARK
            
        # Pie de página oficial
        footer = section.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = f_p.add_run(f"Tombstone Hats MES · Uanify Software Industrial  |  {doc_title}")
        f_run.font.name = "Arial"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = COLOR_MUTED

    with open(md_path, 'r', encoding='utf-8-sig') as f:
        lines = f.readlines()

    in_metadata_card = False
    metadata_lines = []
    in_table = False
    table_rows = []
    in_code = False
    code_lines = []
    is_first_h1 = True

    for line in lines:
        raw = line.rstrip('\r\n')
        stripped = raw.strip()

        # Bloques de Código
        if stripped.startswith('```'):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                tbl = doc.add_table(rows=1, cols=1)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = tbl.cell(0, 0)
                set_cell_background(cell, "0F172A")
                set_cell_margins(cell, top=130, bottom=130, left=180, right=180)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run('\n'.join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(241, 245, 249)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if in_code:
            code_lines.append(raw)
            continue

        # Cajas de Notas / Callouts (> NOTE ...)
        if stripped.startswith('>'):
            in_metadata_card = True
            metadata_lines.append(stripped.lstrip('>').strip())
            continue
        elif in_metadata_card and not stripped.startswith('>'):
            in_metadata_card = False
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, COLOR_BG_CARD)
            set_cell_margins(cell, top=130, bottom=130, left=200, right=200)
            set_card_borders(cell, COLOR_CARD_BORDER)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            
            for m_line in metadata_lines:
                add_formatted_runs(p, m_line, default_color=COLOR_BODY, default_size=9.5)
                p.add_run('\n')
            if p.runs and p.runs[-1].text.endswith('\n'):
                p.runs[-1].text = p.runs[-1].text[:-1]
            doc.add_paragraph().paragraph_format.space_after = Pt(6)
            metadata_lines = []

        # H1 - TÍTULO PRINCIPAL DEL DOCUMENTO
        if stripped.startswith('# '):
            t_text = clean_emojis(stripped[2:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(t_text)
            r.font.name = "Arial"
            r.font.size = Pt(23)
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY_DARK
            is_first_h1 = False
            continue

        # H2 - SECCIÓN PRINCIPAL (ej. "1. Objetivo General", "2. Alcance...")
        if stripped.startswith('## '):
            h_text = clean_emojis(stripped[3:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = "Arial"
            r.font.size = Pt(14)
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY_DARK
            continue

        # H3 - SUBSECCIÓN CON COLOR AZUL UANIFY (#2563EB)
        if stripped.startswith('### '):
            h_text = clean_emojis(stripped[4:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = "Arial"
            r.font.size = Pt(11.5)
            r.bold = True
            r.font.color.rgb = COLOR_BRAND_BLUE
            continue

        # H4 - NIVEL 4 EN SLATE OSCURO
        if stripped.startswith('#### '):
            h_text = clean_emojis(stripped[5:].strip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.bold = True
            r.font.color.rgb = COLOR_PRIMARY_DARK
            continue

        # LÍNEAS DIVISORIAS MARCADAS EN MARKDOWN
        if stripped in ['---', '***', '___']:
            add_horizontal_divider(doc, space_after=8)
            continue

        # TABLAS DE DATOS
        if stripped.startswith('|') and stripped.endswith('|'):
            if '---' in stripped:
                continue
            in_table = True
            cells = [c.strip() for c in stripped[1:-1].split('|')]
            table_rows.append(cells)
            continue
        elif in_table and not (stripped.startswith('|') and stripped.endswith('|')):
            in_table = False
            if table_rows:
                num_cols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(table_rows):
                    for c_idx, val in enumerate(row):
                        cell = tbl.cell(r_idx, c_idx)
                        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(0)
                        cp.paragraph_format.space_after = Pt(0)
                        
                        if r_idx == 0:
                            set_cell_background(cell, COLOR_TH_BG)
                            add_formatted_runs(cp, val, default_color=COLOR_PRIMARY_DARK, default_size=9.5, default_bold=True)
                        else:
                            bg = COLOR_BG_CARD if r_idx % 2 == 0 else "FFFFFF"
                            set_cell_background(cell, bg)
                            add_formatted_runs(cp, val, default_color=COLOR_BODY, default_size=9.0)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            table_rows = []

        # LISTAS Y PÁRRAFOS
        if stripped:
            clean_line = clean_emojis(stripped)
            is_bullet = False
            is_numbered = False
            content_to_parse = clean_line

            if clean_line.startswith('* ') or clean_line.startswith('- '):
                is_bullet = True
                content_to_parse = clean_line[2:].strip()
            elif re.match(r'^\d+\.\s', clean_line):
                is_numbered = True
                m = re.match(r'^\d+\.\s', clean_line)
                num_prefix = m.group(0)
                content_to_parse = clean_line[len(num_prefix):].strip()

            if is_bullet:
                p = doc.add_paragraph(style='List Bullet')
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.25
                add_formatted_runs(p, content_to_parse, default_color=COLOR_BODY, default_size=10)
            elif is_numbered:
                p = doc.add_paragraph(style='List Number')
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.25
                add_formatted_runs(p, content_to_parse, default_color=COLOR_BODY, default_size=10)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.25
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.space_before = Pt(0)
                add_formatted_runs(p, content_to_parse, default_color=COLOR_BODY, default_size=10)

    doc.save(docx_path)
    return docx_path

def upload_styled_doc(file_path, folder_id=None, doc_title=None):
    local_p = Path(file_path)
    if not local_p.exists():
        raise FileNotFoundError(f"No existe el archivo {local_p}")

    title = doc_title or local_p.stem
    temp_docx = local_p.parent / f"{local_p.stem}_temp_styled.docx"
    
    try:
        md_to_docx(str(local_p), str(temp_docx), title)
        
        media = MediaFileUpload(
            str(temp_docx),
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            resumable=True
        )

        drive, docs = get_services()

        q = f"name = '{title}' and mimeType = 'application/vnd.google-apps.document' and trashed = false"
        if folder_id:
            q += f" and '{folder_id}' in parents"

        search = drive.files().list(q=q, fields="files(id, name, webViewLink)").execute()
        files = search.get('files', [])

        if files:
            doc_id = files[0]['id']
            updated_file = drive.files().update(
                fileId=doc_id,
                media_body=media,
                fields='id, name, webViewLink'
            ).execute()
            res = {'status': 'updated', 'id': doc_id, 'name': title, 'link': updated_file.get('webViewLink')}
        else:
            meta = {
                'name': title,
                'mimeType': 'application/vnd.google-apps.document',
                'parents': [folder_id] if folder_id else []
            }
            created_file = drive.files().create(
                body=meta,
                media_body=media,
                fields='id, name, webViewLink'
            ).execute()
            res = {'status': 'created', 'id': created_file['id'], 'name': title, 'link': created_file.get('webViewLink')}
    finally:
        if temp_docx.exists():
            try:
                temp_docx.unlink()
            except Exception:
                pass
                
    return res

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--sync-file', type=str, help='Ruta local de archivo Markdown')
    parser.add_argument('--folder-id', type=str, help='ID de carpeta destino en Drive')
    parser.add_argument('--title', type=str, help='Título del doc')

    args = parser.parse_args()
    if args.sync_file:
        res = upload_styled_doc(args.sync_file, args.folder_id, args.title)
        print(f"Sincronizado: {res['name']} ({res['status']}) -> {res.get('link')}")
